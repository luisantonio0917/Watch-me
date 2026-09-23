from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r'C:\Users\admin\Downloads\Watch Me - Prototipagem\WATCH ME - Relatorio do projeto.docx'
NAVY='12345D'; BLUE='173E70'; TEAL='148F84'; MINT='DDF2EE'; PALE='F2F6F8'; TEXT='33465B'; MUTED='637589'; WHITE='FFFFFF'; RED='A9443B'
doc=Document(); sec=doc.sections[0]
sec.top_margin=Inches(.68); sec.bottom_margin=Inches(.68); sec.left_margin=Inches(.78); sec.right_margin=Inches(.78); sec.header_distance=Inches(.3); sec.footer_distance=Inches(.35)
styles=doc.styles
normal=styles['Normal']; normal.font.name='Aptos'; normal.font.size=Pt(9.4); normal.font.color.rgb=RGBColor.from_string(TEXT); normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.15
for nm,size,color in [('Title',31,NAVY),('Heading 1',21,NAVY),('Heading 2',13,BLUE),('Heading 3',10,TEAL)]:
    st=styles[nm]; st.font.name='Aptos Display'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color); st.paragraph_format.keep_with_next=True
styles['Heading 1'].paragraph_format.space_before=Pt(0); styles['Heading 1'].paragraph_format.space_after=Pt(9); styles['Heading 2'].paragraph_format.space_before=Pt(9); styles['Heading 2'].paragraph_format.space_after=Pt(4)

def shade(cell,fill):
    shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); cell._tc.get_or_add_tcPr().append(shd)
def margins(cell,top=90,start=120,bottom=90,end=120):
    mar=OxmlElement('w:tcMar')
    for name,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        e=OxmlElement('w:'+name); e.set(qn('w:w'),str(val)); e.set(qn('w:type'),'dxa'); mar.append(e)
    cell._tc.get_or_add_tcPr().append(mar)
def ctext(cell,text,bold=False,color=TEXT,size=8.2):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.08
    r=p.add_run(text); r.bold=bold; r.font.name='Aptos'; r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(cell)
def para(text='',style=None):
    p=doc.add_paragraph(style=style); p.add_run(text); return p
def bullet(text):
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(3); p.paragraph_format.left_indent=Inches(.22); p.paragraph_format.first_line_indent=Inches(-.12); p.add_run(text); return p
def label(text):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(5); r=p.add_run(text.upper()); r.bold=True; r.font.size=Pt(7); r.font.color.rgb=RGBColor.from_string(TEAL)
def callout(title,text,fill=MINT,title_color=TEAL):
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; c=t.cell(0,0); shade(c,fill); margins(c,150,190,145,190); c.text=''
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(3); r=p.add_run(title); r.bold=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string(title_color)
    p=c.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.1; r=p.add_run(text); r.font.size=Pt(8.3); r.font.color.rgb=RGBColor.from_string(TEXT)
    para('').paragraph_format.space_after=Pt(0)
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for i,h in enumerate(headers): shade(t.rows[0].cells[i],NAVY); ctext(t.rows[0].cells[i],h,True,WHITE,7.8)
    trPr=t.rows[0]._tr.get_or_add_trPr(); rep=OxmlElement('w:tblHeader'); rep.set(qn('w:val'),'true'); trPr.append(rep)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            if ri%2: shade(cells[i],PALE)
            ctext(cells[i],val,False,TEXT,7.8)
    para('').paragraph_format.space_after=Pt(0)
def pagehead(kicker,title,intro):
    label(kicker); doc.add_heading(title,level=1); para(intro)

header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=header.add_run('WATCH ME   /   DOCUMENTAÇÃO DO PROJETO'); r.font.size=Pt(7); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(MUTED)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('WATCH ME  •  Relatório de andamento  |  '); r.font.size=Pt(7); r.font.color.rgb=RGBColor.from_string(MUTED)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

label('RELATÓRIO DE ANDAMENTO · SETEMBRO 2026')
p=doc.add_paragraph(style='Title'); p.paragraph_format.space_before=Pt(44); p.paragraph_format.space_after=Pt(8); p.add_run('WATCH ME')
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(19); r=p.add_run('Descrição detalhada do projeto e do site'); r.font.size=Pt(19); r.font.color.rgb=RGBColor.from_string(BLUE)
callout('OBJETIVO DESTE DOCUMENTO','Registrar o que foi definido e construído até agora no projeto WATCH ME: a proposta do dispositivo, o conteúdo informativo do site, suas escolhas visuais, a estrutura Node.js e o que ainda precisa ser desenvolvido.',PALE,NAVY)
doc.add_heading('Resumo executivo',level=2)
para('WATCH ME é uma proposta de relógio inteligente voltado a facilitar o pedido de ajuda em situações de risco. A ideia central é permitir que a pessoa usuária acione um alerta no relógio, envie uma notificação a contatos escolhidos previamente e compartilhe a localização disponível naquele momento.')
para('Foi criado um site responsivo em Node.js para apresentar o conceito de forma séria e informativa. A página explica o propósito, descreve o fluxo esperado do alerta, apresenta os componentes técnicos considerados, aborda privacidade e limitações e reserva cinco espaços para fotos do protótipo em um carrossel.')
callout('ESTADO ATUAL','O site é uma apresentação informativa. A aplicação ainda não conecta a um relógio, não coleta localização, não envia mensagens e não possui cadastro de contatos ou integração com serviços de emergência.','FFF2E9',RED)
doc.add_heading('O que este relatório cobre',level=2)
for x in ['A motivação e o escopo funcional definidos para a proposta.','O fluxo de funcionamento e as dependências técnicas previstas.','As seções, a identidade visual e os recursos interativos do site.','Os arquivos Node.js criados, como executar o projeto e as próximas etapas.']: bullet(x)

doc.add_page_break(); pagehead('01 · CONCEITO E FUNCIONAMENTO','A proposta WATCH ME','O conceito descreve um dispositivo vestível de acionamento manual conectado a uma rede pessoal de apoio. O objetivo é transmitir uma informação inicial útil; o relógio não substitui serviços policiais, médicos ou canais oficiais.')
doc.add_heading('Problema que a proposta aborda',level=2)
para('Em uma situação de risco, desbloquear um telefone, abrir um aplicativo e localizar uma função de emergência pode ser difícil. O WATCH ME explora a possibilidade de iniciar um pedido de apoio por meio de uma interação física no pulso e encaminhar a localização a pessoas previamente escolhidas.')
doc.add_heading('Fluxo de uso previsto',level=2)
table(['ETAPA','AÇÃO NO CONCEITO','INFORMAÇÃO / DEPENDÊNCIA'],[('1 · Acionamento','A pessoa pressiona o botão de alerta do relógio.','O controle precisa ser acessível e reduzir acionamentos acidentais.'),('2 · Posicionamento','O dispositivo associa a localização disponível ao alerta.','A disponibilidade e a precisão variam conforme sinal e ambiente.'),('3 · Transmissão','O sistema envia a notificação para os contatos cadastrados.','É necessária conectividade por rede móvel ou por outro dispositivo conectado.'),('4 · Recebimento','Os contatos consultam o alerta e a posição compartilhada.','A resposta depende da disponibilidade e da ação dos destinatários.')],[1.05,2.55,3.25])
doc.add_heading('Elementos técnicos considerados',level=2)
table(['ELEMENTO','PAPEL NO CONCEITO','PONTO A VALIDAR'],[('Botão de alerta','Iniciar manualmente o envio da notificação.','Posição, pressão necessária, discrição, confirmação e prevenção de erro.'),('GPS / GNSS','Estimar a posição do dispositivo para incluí-la no alerta.','Precisão em área aberta, locais internos, obstáculos e tempo de obtenção.'),('Conectividade','Transportar o alerta e a localização aos destinatários.','Cobertura, consumo de energia, dependência de celular e falhas de envio.'),('Lista de contatos','Definir previamente quem recebe as notificações.','Consentimento, atualização, confirmação de recebimento e gestão segura.')],[1.15,2.55,3.15])
callout('LIMITES DE DESEMPENHO','Cobertura de rede, bateria, sinal de posicionamento, configuração do aparelho e integração entre componentes afetam a operação. O site apresenta o comportamento pretendido, mas não comprova desempenho de um produto.',PALE,NAVY)

doc.add_page_break(); pagehead('02 · SITE INFORMATIVO','Conteúdo e organização da página','O site apresenta a proposta com uma hierarquia direta: começa pelo propósito, explica o fluxo e os componentes e termina com privacidade e limites. A navegação superior leva às seções principais.')
table(['SEÇÃO','CONTEÚDO APRESENTADO'],[('Abertura','Resumo da proposta, chamada para explorar o funcionamento e aviso de que se trata de conceito em desenvolvimento.'),('O projeto','Contextualiza a necessidade, explica o objetivo do relógio e resume acionamento, contatos e localização.'),('Fluxo de funcionamento','Apresenta em quatro etapas o acionamento, posicionamento, transmissão e recebimento do alerta.'),('Componentes e requisitos','Descreve botão, posicionamento, conectividade e gestão de contatos; registra itens a validar.'),('Galeria','Carrossel com cinco espaços provisórios para imagens do dispositivo e do processo de prototipagem.'),('Privacidade e limites','Explica consentimento, transparência e proteção dos dados de localização; informa o estado conceitual e limitações.'),('Encerramento','Retoma a finalidade do projeto e oferece um caminho de volta ao início.')],[1.7,5.15])
doc.add_heading('Direção visual aplicada',level=2)
para('A apresentação visual foi refeita seguindo a referência fornecida: cabeçalho branco e limpo, área de abertura azul-escura, detalhes em turquesa e texto em tons neutros de alto contraste. O layout usa seções bem delimitadas, títulos objetivos, cartões informativos e espaçamento regular. A intenção é aproximar o projeto de um portal informativo profissional, em vez de uma página promocional lúdica.')
table(['DECISÃO','APLICAÇÃO NO SITE'],[('Paleta','Azul para confiança e estrutura; turquesa para ações e destaques; branco e cinzas claros para leitura.'),('Tipografia','Famílias sem serifa, com títulos de maior peso e texto corrido simples.'),('Hierarquia','Números de seção, subtítulos e etiquetas identificam a função de cada bloco.'),('Responsividade','O layout reorganiza a navegação, os cartões e as colunas em telas menores.'),('Acessibilidade básica','Idioma pt-BR, link para pular ao conteúdo, rótulos acessíveis nos controles e suporte a movimento reduzido.')],[1.45,5.4])
doc.add_heading('Carrossel de imagens',level=2)
para('Foram preparados cinco espaços visuais identificados para receber fotografias do relógio, interface, componentes, desenvolvimento e contexto de uso. O carrossel tem botões de anterior e próximo, indicadores clicáveis e contador de posição. As imagens atuais são apenas espaços reservados; não foram fornecidas fotos reais do protótipo.')

doc.add_page_break(); pagehead('03 · IMPLEMENTAÇÃO','Estrutura Node.js e arquivos','A aplicação usa um servidor HTTP nativo do Node.js e arquivos estáticos locais. Essa estrutura mantém a apresentação simples e executável sem instalar dependências de terceiros.')
table(['ARQUIVO','RESPONSABILIDADE'],[('package.json','Define metadados básicos do projeto e o comando npm start.'),('server.js','Inicia o servidor HTTP, serve arquivos de public, define tipos MIME básicos e usa a porta indicada por PORT ou a porta 3000.'),('public/index.html','Contém o conteúdo em português, a estrutura semântica da página, as seções do site, o desenho conceitual do relógio e os cinco espaços de foto.'),('public/styles.css','Define cores, tipografia, layout, aparência do relógio ilustrativo, cartões, responsividade e regras para movimento reduzido.'),('public/app.js','Controla o carrossel: navegação, pontos, contador, rotação automática e pausa quando o cursor está sobre o carrossel ou o foco está dentro dele.')],[1.65,5.2])
doc.add_heading('Como executar localmente',level=2)
callout('COMANDO','Abra um terminal na pasta do projeto e execute:  npm start\nDepois, abra http://localhost:3000 no navegador.',MINT,TEAL)
para('A porta pode ser alterada por meio da variável de ambiente PORT. O servidor lê apenas arquivos dentro da pasta public e retorna erro 404 quando o arquivo solicitado não existe.')
doc.add_heading('O que a implementação não contém',level=2)
for x in ['Não existe API para criar, editar ou armazenar uma lista de contatos de emergência.','Não há integração com GPS, GNSS, Bluetooth, modem celular ou relógio físico.','Nenhum alerta é enviado por SMS, ligação, aplicativo ou outro canal.','A localização exibida na ilustração e os estados do relógio são elementos demonstrativos, não dados ao vivo.','Não há autenticação, banco de dados, painel de usuário, mapa ou monitoramento em tempo real.']: bullet(x)
callout('DISTINÇÃO IMPORTANTE','A expressão “em tempo real” descreve a intenção da funcionalidade do produto futuro. No site atual, não há localização em tempo real nem transmissão de alertas.','FFF2E9',RED)

doc.add_page_break(); pagehead('04 · ESTADO E PRÓXIMAS ETAPAS','O que foi feito e o que falta','O trabalho realizado até agora estabeleceu a apresentação do conceito e criou a base visual e técnica de uma página informativa. A etapa seguinte é decidir e validar o funcionamento do produto, antes de comunicar capacidades como funções concluídas.')
doc.add_heading('Concluído até o momento',level=2)
for x in ['Definição do nome WATCH ME e do conceito de relógio inteligente com botão de alerta, contatos predefinidos e compartilhamento de localização.','Criação de um site de apresentação em Node.js, HTML, CSS e JavaScript sem dependências externas.','Reformulação da identidade visual para aproximar o resultado da referência: azul, turquesa, fundo claro e hierarquia informativa.','Inclusão de conteúdo sobre fluxo de uso, componentes, condições técnicas, privacidade e limitações.','Criação do carrossel funcional com cinco espaços preparados para receber fotografias reais.']: bullet(x)
doc.add_heading('Próximas etapas recomendadas',level=2)
table(['ETAPA','TRABALHO A REALIZAR','RESULTADO ESPERADO'],[('1 · Definir requisitos','Escolher se o relógio se conecta diretamente à rede móvel ou depende de um telefone; definir destinatários, conteúdo do alerta e comportamento em caso de falha.','Escopo técnico e de uso verificável.'),('2 · Criar protótipo físico','Selecionar placa, botão, fonte de energia, módulo de posicionamento e meio de comunicação.','Dispositivo experimental que permita testes controlados.'),('3 · Implementar serviço','Criar backend, interface de configuração e método de entrega de notificações, com autenticação e registro técnico apropriados.','Fluxo integrado, sem expor contatos ou localização indevidamente.'),('4 · Proteger dados','Definir consentimento, acesso, armazenamento, retenção, exclusão e tratamento de incidentes.','Regras documentadas de privacidade e segurança.'),('5 · Testar com segurança','Medir precisão, tempo de envio, autonomia, cobertura, falhas, uso acidental e acessibilidade.','Evidências sobre limites reais e melhorias necessárias.'),('6 · Atualizar o site','Adicionar fotos autênticas, resultados medidos e status de cada função, sem apresentar planos como capacidades disponíveis.','Apresentação fiel ao estado comprovado do projeto.')],[1.2,3.4,2.25])
doc.add_heading('Síntese',level=2)
para('Até este momento, WATCH ME conta com uma proposta documentada e um site informativo responsivo. O material explica o valor pretendido e registra dependências e riscos, mas o produto eletrônico e os serviços de localização e envio de alerta continuam como trabalho futuro. Fotos reais, testes e decisões de arquitetura ainda serão necessários para avançar de conceito para protótipo funcional.')
label('FIM DO RELATÓRIO · WATCH ME')
doc.core_properties.title='WATCH ME - Relatório do projeto'; doc.core_properties.subject='Registro detalhado do conceito, site e estado de implementação do projeto WATCH ME'; doc.core_properties.author='WATCH ME'
doc.save(OUT); print(OUT)
